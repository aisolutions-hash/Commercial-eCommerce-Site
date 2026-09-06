"""Regenerate kalisoftai-datahub/features_progress_tracker.xlsx (run: python tools/make_tracker.py)."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

OUT = Path(__file__).resolve().parents[1] / "features_progress_tracker.xlsx"

FEATURES = [
    # (feature, status, branch, owner, notes)
    ("Sales insights dashboard", "shipped", "main", "Kalika", "product analytics"),
    ("Sales contacts + leads", "shipped", "main", "Kalika", "CRUD + filters"),
    ("AI pipeline (data prep->release)", "shipped", "main", "Kalika", "gemma env/demo mode"),
    ("Sync to kalika_enterprises", "shipped", "main", "Kalika", "UI button + scheduler"),
    ("Chat-assisted review", "shipped", "main", "Kalika", "prompts from /prompts"),
    ("Sales/marketing views", "shipped", "main", "Kalika", "created at startup"),
    ("Auth role matrix", "future", "feat/auth-roles", "Kalika", "RBAC per role"),
    ("RL bandit optimization", "future", "feat/rl-bandit", "Kalika", "epsilon-greedy prompts"),
    ("Cloud Function sync", "future", "feat/gcp-sync", "Kalika", "Cloud Scheduler + Function"),
    ("BigQuery warehouse sync", "future", "feat/bigquery", "Kalika", "partition + batch id"),
]

STATUS_FILL = {
    "shipped": PatternFill("solid", fgColor="C6EFCE"),
    "future": PatternFill("solid", fgColor="FFEB9C"),
}
HDR_FILL = PatternFill("solid", fgColor="4472C4")
HDR_FONT = Font(color="FFFFFF", bold=True)

wb = Workbook()

ws = wb.active
ws.title = "Features"
ws.append(["Feature", "Status", "Branch", "Owner", "Notes", "Progress %"])
for cell in ws[1]:
    cell.fill = HDR_FILL
    cell.font = HDR_FONT
    cell.alignment = Alignment(horizontal="center")
progress = 60  # 6 of 10 shipped
for i, row in enumerate(FEATURES, start=2):
    ws.append([*row, 100 if row[1] == "shipped" else 0])
    ws.cell(i, 2).fill = STATUS_FILL[row[1]]
ws.column_dimensions["A"].width = 34
for col in "BCDE":
    ws.column_dimensions[col].width = 22
ws.column_dimensions["F"].width = 12

ws2 = wb.create_sheet("Branches")
ws2.append(["Branch", "Purpose", "Status"])
for c in ws2[1]:
    c.fill = HDR_FILL
    c.font = HDR_FONT
branches = [
    ("main", "working feature line", "active"),
    ("feat/auth-roles", "role-based access for datahub", "planned"),
    ("feat/rl-bandit", "continuous learning loop", "planned"),
    ("feat/gcp-sync", "cloud function data transfer", "planned"),
    ("feat/bigquery", "warehouse mirroring", "planned"),
]
for b in branches:
    ws2.append(b)
ws2.column_dimensions["A"].width = 18
ws2.column_dimensions["B"].width = 40
ws2.column_dimensions["C"].width = 12

ws3 = wb.create_sheet("Progress Log")
ws3.append(["Date", "Milestone", "Status"])
for c in ws3[1]:
    c.fill = HDR_FILL
    c.font = HDR_FONT
log = [
    ("2026-09-06", "Sales insights dashboard live", "done"),
    ("2026-09-06", "Sales contacts + leads live", "done"),
    ("2026-09-06", "AI pipeline + guardrails live", "done"),
    ("2026-09-06", "kalika_enterprises sync + schedule live", "done"),
    ("2026-09-06", "Chat-assisted review live", "done"),
    ("next", "Auth role matrix", "todo"),
    ("next", "RL bandit loop", "todo"),
    ("next", "GCP Cloud Function sync", "todo"),
]
for r in log:
    ws3.append(r)
ws3.column_dimensions["A"].width = 14
ws3.column_dimensions["B"].width = 44
ws3.column_dimensions["C"].width = 12

wb.save(OUT)
print(f"written: {OUT} (features progress {progress}%)")